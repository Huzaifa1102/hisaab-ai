"""
Generate professional PPTX and PDF presentation for Hisaab AI hackathon submission:
TeamName_Domain_LeaderName -> HisaabAI_FinTech_Huzaifa
"""

import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_DIR = Path(r"C:\Users\Hp\.gemini\antigravity-ide\scratch\hisaab_ai")
PPTX_FILE = OUTPUT_DIR / "HisaabAI_FinTech_Huzaifa.pptx"
PDF_FILE = OUTPUT_DIR / "HisaabAI_FinTech_Huzaifa.pdf"

# Color Palette: Modern Fintech Dark & Emerald Green
NAVY = RGBColor(15, 23, 42)        # Slate 900
DARK_TEAL = RGBColor(6, 78, 59)    # Emerald 900
EMERALD = RGBColor(16, 185, 129)   # Emerald 500
LIGHT_BG = RGBColor(248, 250, 252) # Slate 50
TEXT_DARK = RGBColor(30, 41, 59)   # Slate 800
TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(245, 158, 11)      # Amber 500
CARD_BG = RGBColor(255, 255, 255)
BORDER_COL = RGBColor(226, 232, 240)

prs = Presentation()
prs.slide_width = Inches(13.333) # 16:9 Widescreen
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_text="HISAAB AI • PITCH DECK"):
    # Category tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = EMERALD
    
    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = NAVY

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    return shape

# ==========================================
# SLIDE 1: Title & Hook
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = NAVY
bg1.line.fill.background()

# Tag badge
badge = create_card(s1, Inches(1.0), Inches(1.2), Inches(3.6), Inches(0.45), bg_color=DARK_TEAL, border_color=EMERALD)
tf_b = badge.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "🏆 BUILD WITH AI HACKATHON 2026"
p_b.alignment = PP_ALIGN.CENTER
p_b.font.size = Pt(12)
p_b.font.bold = True
p_b.font.color.rgb = WHITE

# Main Title
title_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(2.2))
tf1 = title_box.text_frame
tf1.word_wrap = True
p1 = tf1.paragraphs[0]
p1.text = "Hisaab AI (حساب)"
p1.font.size = Pt(54)
p1.font.bold = True
p1.font.color.rgb = WHITE

p2 = tf1.add_paragraph()
p2.text = "Multimodal AI Bookkeeper & Cash Drawer Assistant for Pakistani Kiryana Stores"
p2.font.size = Pt(22)
p2.font.color.rgb = EMERALD

p3 = tf1.add_paragraph()
p3.text = "Empowering 2M+ micro-retailers with Computer Vision, Urdu Localization, WhatsApp Reminders & On-Device Resilience."
p3.font.size = Pt(15)
p3.font.color.rgb = RGBColor(203, 213, 225)

# Meta info card
meta_card = create_card(s1, Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.5), bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
tf_m = meta_card.text_frame
tf_m.word_wrap = True

pm1 = tf_m.paragraphs[0]
pm1.text = "👤 Team: HisaabAI   |   Domain: FinTech / Generative AI   |   Team Lead: Huzaifa"
pm1.font.size = Pt(15)
pm1.font.bold = True
pm1.font.color.rgb = WHITE

pm2 = tf_m.add_paragraph()
pm2.text = "🚀 Tech: Google Gemini Flash Multimodal, Groq Whisper-v3, Streamlit, SQLite, WhatsApp Cloud / Twilio API"
pm2.font.size = Pt(13)
pm2.font.color.rgb = RGBColor(148, 163, 184)

pm3 = tf_m.add_paragraph()
pm3.text = "🔗 GitHub: github.com/Huzaifa1102/hisaab-ai   •   Status: Production Deployed"
pm3.font.size = Pt(13)
pm3.font.color.rgb = GOLD

# ==========================================
# SLIDE 2: The Problem
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "The Problem: The Fragile Reality of Kiryana Retail", "Market Pain Points")

problems = [
    ("1. The Manual Khata Trap", "Over 2 Million neighborhood grocery stores in Pakistan run solely on handwritten paper registers. Illegible handwriting and torn notebooks lead to untracked customer debt and forgotten loans."),
    ("2. Dead Capital & Stockout Losses", "Shopkeepers lack real-time stock ledgers. Capital gets locked in slow-moving goods while high-demand items (Atta, Cheeni, Ghee) run out, causing direct revenue losses."),
    ("3. Cash Drawer Leakage (Galla)", "At the end of a 14-hour shift, cash in the physical drawer rarely matches transactions. Without reconciliation tools, discrepancies, miscalculations, and theft go unnoticed."),
    ("4. Digital Literacy Barrier", "Conventional POS software is English-first, desktop-bound, expensive, and complex. Kiryana owners need native Urdu script, voice support, and zero learning curve.")
]

for idx, (title, desc) in enumerate(problems):
    row = idx // 2
    col = idx % 2
    x = Inches(0.8 + col * 5.9)
    y = Inches(1.7 + row * 2.6)
    
    card = create_card(s2, x, y, Inches(5.6), Inches(2.3))
    tf = card.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    p_body = tf.add_paragraph()
    p_body.text = desc
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 3: The Solution
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "The Solution: Hisaab AI Architecture", "Autonomous Bookkeeper")

solutions = [
    ("📷 Smart Vision Scanner", "Instantly ingests handwritten Urdu receipts, wholesaler invoices, and digital utility bill PDFs (IESCO, SNGPL). Extracts line items, vendor name, amount, and due dates."),
    ("🎙️ Roman Urdu Voice Khata", "Powered by Groq Whisper-large-v3. Shopkeepers speak naturally: 'Faisal ko 500 ka ghee udhaar diya' — transcribed and parsed directly into financial entries."),
    ("💬 Automated WhatsApp Reminders", "Generates polite, bilingual payment reminders with exact balances. Shopkeepers dispatch reminders via WhatsApp Web or automated Cloud API with one click."),
    ("💵 Galla Hisaab (Reconciliation)", "Automated end-of-day drawer balancing. Reconciles opening cash + cash sales + repayments vs supplier expenses, featuring an intuitive currency note counter.")
]

for idx, (title, desc) in enumerate(solutions):
    row = idx // 2
    col = idx % 2
    x = Inches(0.8 + col * 5.9)
    y = Inches(1.7 + row * 2.6)
    
    card = create_card(s3, x, y, Inches(5.6), Inches(2.3), border_color=EMERALD)
    tf = card.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = DARK_TEAL
    
    p_body = tf.add_paragraph()
    p_body.text = desc
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 4: Core Enterprise Features
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Enterprise-Ready Features Built for Hackathon", "Product Depth")

features = [
    ("🇵🇰 100% Urdu Localization", "Instant one-click toggle between English and Urdu Nastaliq script with full RTL alignment across all screens."),
    ("📦 Inventory Stock Ledger", "Real-time stock valuation (Cost vs Retail), projected profit margins, safety reorder alerts, and auto-restock from invoice scans."),
    ("🏪 Multi-Shop & 4-Digit PIN", "Multi-tenancy profile switcher (e.g. Madina Store, Bismillah Kiryana) with PIN protection for ledger security."),
    ("📊 Financial Intelligence", "Real-time KPIs, monthly expense trends, unpaid bill trackers, and 'Ask AI' Kiryana financial advisor.")
]

for idx, (title, desc) in enumerate(features):
    row = idx // 2
    col = idx % 2
    x = Inches(0.8 + col * 5.9)
    y = Inches(1.7 + row * 2.6)
    
    card = create_card(s4, x, y, Inches(5.6), Inches(2.3))
    tf = card.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    p_body = tf.add_paragraph()
    p_body.text = desc
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 5: Resilience & Offline Architecture
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Built for Pakistan's Reality: 100% Fault Tolerance", "Technical Architecture")

arch_points = [
    ("🛡️ Multi-Key Failover Pool", "Eliminates API 429 quota exhaustion. Auto-rotates across primary and backup Gemini API keys seamlessly in the background."),
    ("⚡ On-Device Heuristic Fallback", "Unseen images and PDFs are parsed locally without internet access using smart heuristic pattern matching and local OCR."),
    ("🔒 Zero Key Leakage Security", "All API keys are masked with password inputs and strictly stored in local .env, never exposed or committed."),
    ("💾 Local-First Persistence", "Self-contained SQLite database engine with zero external database dependencies or monthly server costs.")
]

for idx, (title, desc) in enumerate(arch_points):
    row = idx // 2
    col = idx % 2
    x = Inches(0.8 + col * 5.9)
    y = Inches(1.7 + row * 2.6)
    
    card = create_card(s5, x, y, Inches(5.6), Inches(2.3), border_color=NAVY)
    tf = card.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    p_body = tf.add_paragraph()
    p_body.text = desc
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 6: Impact & Live Links
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
bg6 = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg6.fill.solid()
bg6.fill.fore_color.rgb = NAVY
bg6.line.fill.background()

# Title
t_box6 = s6.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.5))
tf6 = t_box6.text_frame
tf6.word_wrap = True
p_head = tf6.paragraphs[0]
p_head.text = "Transforming Informal Retail in Pakistan"
p_head.font.size = Pt(36)
p_head.font.bold = True
p_head.font.color.rgb = WHITE

p_sub = tf6.add_paragraph()
p_sub.text = "Bridging the gap between traditional Kiryana shops and digital financial inclusion."
p_sub.font.size = Pt(18)
p_sub.font.color.rgb = EMERALD

# Summary Box
sbox = create_card(s6, Inches(1.0), Inches(2.5), Inches(11.3), Inches(4.2), bg_color=RGBColor(30, 41, 59), border_color=EMERALD)
tf_s = sbox.text_frame
tf_s.word_wrap = True

ps1 = tf_s.paragraphs[0]
ps1.text = "🚀 PROJECT SUMMARY & VERIFIED HACKATHON DELIVERABLES"
ps1.font.size = Pt(18)
ps1.font.bold = True
ps1.font.color.rgb = WHITE

items = [
    ("Live Demonstration:", "Hosted on Streamlit Community Cloud with instant demo-seed utilities"),
    ("Code Repository:", "https://github.com/Huzaifa1102/hisaab-ai (100% Passing Automated Tests)"),
    ("Localization:", "Full English & Urdu Nastaliq interfaces tailored for local shopkeepers"),
    ("Multimodal AI Pipeline:", "Google Gemini Flash Vision + Groq Whisper Speech Recognition"),
    ("Resilience Engine:", "Multi-Key API Pool + Offline Document Parsing Fallback"),
    ("Team Lead:", "Huzaifa | Team: HisaabAI | Domain: FinTech / AI")
]

for label, val in items:
    p_item = tf_s.add_paragraph()
    p_item.text = f"• {label} {val}"
    p_item.font.size = Pt(14)
    p_item.font.color.rgb = RGBColor(226, 232, 240)

prs.save(str(PPTX_FILE))
print(f"PPTX successfully created at: {PPTX_FILE}")

# Now generate a companion high-quality PDF using reportlab
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen import canvas

c = canvas.Canvas(str(PDF_FILE), pagesize=landscape(letter))
width, height = landscape(letter)

def draw_pdf_slide(title, subtitle, bullets, is_dark=False):
    if is_dark:
        c.setFillColorRGB(0.06, 0.09, 0.16) # Navy
        c.rect(0, 0, width, height, fill=True, stroke=False)
        c.setFillColorRGB(0.06, 0.72, 0.50) # Emerald
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, height - 50, "BUILD WITH AI HACKATHON 2026")
        c.setFillColorRGB(1, 1, 1)
        c.setFont("Helvetica-Bold", 26)
        c.drawString(50, height - 90, title)
        c.setFillColorRGB(0.06, 0.72, 0.50)
        c.setFont("Helvetica", 15)
        c.drawString(50, height - 120, subtitle)
        c.setFillColorRGB(0.9, 0.9, 0.95)
        c.setFont("Helvetica", 13)
        y = height - 170
        for b in bullets:
            c.drawString(70, y, f"•  {b}")
            y -= 32
    else:
        c.setFillColorRGB(0.98, 0.98, 0.99)
        c.rect(0, 0, width, height, fill=True, stroke=False)
        c.setFillColorRGB(0.06, 0.72, 0.50)
        c.setFont("Helvetica-Bold", 11)
        c.drawString(50, height - 50, "HISAAB AI • PITCH DECK")
        c.setFillColorRGB(0.06, 0.09, 0.16)
        c.setFont("Helvetica-Bold", 24)
        c.drawString(50, height - 85, title)
        c.setFillColorRGB(0.2, 0.3, 0.4)
        c.setFont("Helvetica", 13)
        c.drawString(50, height - 110, subtitle)
        y = height - 155
        for b in bullets:
            c.setFillColorRGB(0.06, 0.09, 0.16)
            c.setFont("Helvetica-Bold", 13)
            c.drawString(60, y, b[0])
            c.setFillColorRGB(0.3, 0.35, 0.45)
            c.setFont("Helvetica", 11)
            c.drawString(75, y - 18, b[1])
            y -= 48
    c.showPage()

# PDF Slide 1
draw_pdf_slide(
    "Hisaab AI (حساب)",
    "Multimodal AI Bookkeeper & Cash Drawer Assistant for Pakistani Kiryana Stores",
    [
        "Team: HisaabAI | Domain: FinTech / Generative AI | Leader: Huzaifa",
        "Empowering 2M+ micro-retailers across Pakistan with AI-powered financial management.",
        "Tech Stack: Google Gemini Flash Multimodal, Groq Whisper-v3, Streamlit, SQLite.",
        "GitHub: github.com/Huzaifa1102/hisaab-ai • Fully Deployed & Tested"
    ],
    is_dark=True
)

# PDF Slide 2
draw_pdf_slide(
    "The Problem: The Fragile Reality of Kiryana Retail",
    "Why 2M+ small retail shops in Pakistan struggle with profitability",
    [
        ("1. The Manual Khata Trap", "Paper notebooks get lost, torn, or smudged, leaving thousands of rupees in uncollected customer debt."),
        ("2. Dead Capital & Stockouts", "No inventory visibility leads to locked cash in slow items while fast-moving commodities run out."),
        ("3. Cash Drawer Leakage (Galla)", "Discrepancies between cash in the till and actual sales go unnoticed without daily reconciliation."),
        ("4. The Digital Literacy Barrier", "Conventional POS software is in English, expensive, desktop-bound, and too complex for shopkeepers.")
    ]
)

# PDF Slide 3
draw_pdf_slide(
    "The Solution: Hisaab AI Architecture",
    "Tailored multimodal intelligence designed specifically for local shopkeeper workflows",
    [
        ("Smart Vision Scanner", "Extracts line items, totals, and due dates from handwritten receipts, invoices, and PDF utility bills."),
        ("Roman Urdu Voice Khata", "Groq Whisper-large-v3 transcribes spoken notes ('Faisal ko 500 ka ghee udhaar diya') into ledger entries."),
        ("Automated WhatsApp Reminders", "Generates polite bilingual payment reminders dispatched via WhatsApp Web or Cloud API with one click."),
        ("Galla Hisaab (Reconciliation)", "Balances daily physical cash drawer against sales, repayments, and expenses with a currency counter.")
    ]
)

# PDF Slide 4
draw_pdf_slide(
    "Core Enterprise Features",
    "Production-grade depth built for real-world Kiryana store environments",
    [
        ("100% Urdu Localization", "Instant toggle between English and authentic Urdu Nastaliq with complete RTL layout support."),
        ("Commodity Inventory Ledger", "Tracks stock levels, cost vs retail valuation, profit margins, and auto-restocks from scanned bills."),
        ("Multi-Store & Security PIN", "Multi-tenancy profile switcher with 4-digit PIN lock safeguarding sensitive ledger resets."),
        ("Financial Intelligence & Ask AI", "Visual expense charts, unpaid bill alerts, and a conversational AI Kiryana financial advisor.")
    ]
)

# PDF Slide 5
draw_pdf_slide(
    "Resilience & 100% Fault Tolerance",
    "Engineered to withstand internet outages and API rate limits",
    [
        ("Multi-Key Failover Pool", "Automatically rotates through backup Gemini API keys if quota 429 limits are reached."),
        ("On-Device Heuristic Fallback", "Unseen images and PDFs are parsed locally on-device without internet access."),
        ("Zero Key Leakage Security", "API keys are masked in password fields, saved in local .env, and never committed to GitHub."),
        ("Local-First Persistence", "Self-contained SQLite architecture with zero external database dependencies or recurring fees.")
    ]
)

# PDF Slide 6
draw_pdf_slide(
    "Transforming Informal Retail in Pakistan",
    "Bridging traditional Kiryana commerce with digital financial inclusion",
    [
        "Live App: Hosted and accessible for judges with 1-click demo ledger seeding",
        "Source Code: Publicly available at https://github.com/Huzaifa1102/hisaab-ai",
        "Complete Localization: Authentic Urdu Nastaliq & English",
        "Next Steps: Micro-lending credit scoring for unbanked shopkeepers",
        "Team Lead: Huzaifa | Team: HisaabAI | Domain: FinTech"
    ],
    is_dark=True
)

c.save()
print(f"PDF successfully created at: {PDF_FILE}")
